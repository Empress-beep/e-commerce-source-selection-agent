"""待审核 Excel 批量导入：业务人工填写 -> 回写 MySQL -> 归档 JSON/Excel。

用法
----
    python src/selector/excel_audit_import.py --file "./output/excel/20261008_113918_442_pending_audit.xlsx"

流程
----
1. 读取 ``ai_select_product`` 产出的 ``*_pending_audit.xlsx``（openpyxl），
   必须列：``id``（数据库主键，回写匹配键）、``audit_status``、
   ``audit_operator``、``audit_comment``；
2. 逐行处理：
   * 完全空白行直接跳过；
   * ``audit_status`` 为空 -> 打印“该行id=xxx尚未填写审核状态，跳过处理”，不更新库；
   * 否则 **复用** :func:`src.selector.manual_audit.manual_audit_product`
     （状态机校验的唯一实现：流转规则 / reject必填备注 / 终态禁止重复修改，
     本文件禁止另写一套校验）；
   * 单行异常只记录（行号/id/报错），**绝不中断整个脚本**，全部失败收集完；
3. 遍历结束后：按本次处理过的 id 从 MySQL 查最新完整数据 -> 归档
   ``output/json_dump/{ts}_finished_audit.json``（AI原始字段+人工审核字段）
   与 ``output/excel/{ts}_finished_audit.xlsx``（数据库真实最新状态）；
4. 控制台统计：成功更新N条 / 跳过N条未填写 / 失败N条 + 全部失败明细。

约束
----
* 只更新 ``audit_status`` / ``audit_operator`` / ``audit_comment`` /
  ``audit_time`` 4 个审核字段（由 manual_audit_product 保证），AI 原始
  评估字段（composite_score / recommend_level / risk_list_json / reason /
  select_condition_json 等）只读不改；
* 按主键 ``id`` 单行 UPDATE，禁止 pandas to_sql 整表覆盖；
* 数据库连接统一走 :mod:`src.common.database`。
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

# sys.path 引导：把项目根（src/ 的上一级）加进搜索路径，
# 支持直接脚本启动 `python src/selector/excel_audit_import.py --file ...`
# 与 `python -m src.selector.excel_audit_import` 两种方式
_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.common import database  # noqa: E402
from src.selector.manual_audit import manual_audit_product  # noqa: E402
from src.settings import OUTPUT_DIR  # noqa: E402

# Excel 必须提供的列（表头按去空格后的字符串精确匹配）
REQUIRED_COLUMNS = ("id", "audit_status", "audit_operator", "audit_comment")
TABLE = "product_select_result"


def _cell_text(value) -> str:
    """单元格值 -> 去首尾空白的字符串（None -> ""）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_row_id(value) -> int:
    """单元格主键 -> int（Excel 数字可能读成 113.0 这类浮点）。"""
    text = _cell_text(value)
    if not text:
        raise ValueError("id 为空")
    try:
        return int(text)
    except ValueError:
        return int(float(text))  # 兼容 "113.0"


def _read_header(ws) -> dict[str, int]:
    """读表头行 -> {列名: 1-based 列号}；缺列直接报错退出。"""
    header = {}
    for col in range(1, ws.max_column + 1):
        name = _cell_text(ws.cell(row=1, column=col).value)
        if name:
            header[name] = col
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        print(f"[ERROR] Excel 缺少必需列: {missing}（实际表头: {sorted(header)}）")
        sys.exit(2)
    return header


def _archive(processed_ids: list[int], ts: str, source_file: Path, stats: dict) -> None:
    """按处理过的 id 查库最新全量数据 -> 归档 JSON + 归档 Excel。"""
    placeholders = ", ".join(["%s"] * len(processed_ids))
    rows = database.fetch_all(
        TABLE,
        where=f"id IN ({placeholders}) ORDER BY id",
        params=processed_ids,
    )
    if not rows:
        print("[归档] 未查询到待归档数据，跳过归档输出")
        return

    json_dir = OUTPUT_DIR / "json_dump"
    excel_dir = OUTPUT_DIR / "excel"
    json_dir.mkdir(parents=True, exist_ok=True)
    excel_dir.mkdir(parents=True, exist_ok=True)

    # ① 归档 JSON：AI 原始选品全部字段 + 人工审核全部字段（SELECT * 全列）
    json_path = json_dir / f"{ts}_finished_audit.json"
    payload = {
        "archived_at": ts,
        "source_file": str(source_file),
        "count": len(rows),
        "stats": stats,
        "rows": rows,  # id/product_id/.../select_condition_json + audit_* 全字段
    }
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    # ② 归档 Excel：数据库真实最新状态（业务归档留存）
    columns = list(rows[0].keys())
    xlsx_path = excel_dir / f"{ts}_finished_audit.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "审核完成归档"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])
    for idx, col in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = (
            14 if col.startswith("audit_") else max(12, min(40, len(col) + 4))
        )
    wb.save(xlsx_path)

    print(f"[归档] JSON: output/json_dump/{json_path.name}（{len(rows)} 行）")
    print(f"[归档] Excel: output/excel/{xlsx_path.name}（{len(rows)} 行）")


def import_audit_excel(file_path: str) -> int:
    """读取人工填写完成的待审核 Excel 并批量回写 + 归档，返回进程退出码。"""
    path = Path(file_path)
    if not path.exists():
        print(f"[ERROR] 文件不存在: {path}")
        return 2

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except OSError as exc:
        print(f"[ERROR] Excel 打开失败: {exc}")
        return 2
    ws = wb.active
    header = _read_header(ws)
    col_id = header["id"]
    col_status = header["audit_status"]
    col_operator = header["audit_operator"]
    col_comment = header["audit_comment"]

    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    success_ids: list[int] = []      # 成功更新的 id（归档用）
    failures: list[tuple[int, str, str]] = []  # (行号, id, 报错)
    skipped = 0
    blank = 0
    max_row = ws.max_row

    print(f"[导入] 源文件: {path}")
    for row_no in range(2, max_row + 1):
        raw_id = ws.cell(row=row_no, column=col_id).value
        status = _cell_text(ws.cell(row=row_no, column=col_status).value)
        # ---- 完全空白行：直接跳过 ----
        if not any(
            _cell_text(ws.cell(row=row_no, column=c).value)
            for c in header.values()
        ):
            blank += 1
            continue
        # ---- 未填写审核状态：只记日志，不更新库 ----
        if not status:
            print(f"该行id={_cell_text(raw_id)}尚未填写审核状态，跳过处理")
            skipped += 1
            continue
        # ---- 填写了状态：复用 manual_audit_product（唯一校验实现） ----
        operator = _cell_text(ws.cell(row=row_no, column=col_operator).value)
        comment = _cell_text(ws.cell(row=row_no, column=col_comment).value)
        try:
            select_result_id = _to_row_id(raw_id)
            manual_audit_product(select_result_id, status, operator, comment)
            success_ids.append(select_result_id)
        except Exception as exc:  # noqa: BLE001 - 单行失败不中断整个脚本
            failures.append((row_no, _cell_text(raw_id) or "-", str(exc)))
    wb.close()
    if blank:
        print(f"[导入] 忽略完全空白行 {blank} 行")

    # ---- ③ 归档：数据库真实最新状态（AI原始字段 + 人工审核字段） ----
    stats = {
        "success": len(success_ids),
        "skipped": skipped,
        "failed": len(failures),
        "total_rows": max_row - 1,
    }
    if success_ids:
        _archive(success_ids, ts, path, stats)
    else:
        print("[归档] 本次无成功更新记录，跳过归档输出")

    # ---- ④ 控制台统计汇总 ----
    print("=" * 46)
    print(f"[汇总] 成功更新 {len(success_ids)} 条")
    print(f"[汇总] 跳过 {skipped} 条未填写审核状态")
    print(f"[汇总] 失败 {len(failures)} 条")
    if failures:
        print("---- 失败明细 ----")
        for row_no, id_text, err in failures:
            print(f"  第{row_no}行 id={id_text}: {err}")
    print("=" * 46)
    return 0


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(
        description="读取人工填写完成的待审核Excel，批量回写 MySQL product_select_result 并归档"
    )
    parser.add_argument(
        "--file", required=True,
        help='待审核Excel路径，如 "./output/excel/20261008_113918_442_pending_audit.xlsx"',
    )
    args = parser.parse_args()
    return import_audit_excel(args.file)


if __name__ == "__main__":
    sys.exit(main())
