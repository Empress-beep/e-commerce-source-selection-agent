"""P1 素材(文案)生成：``generate_goods_material`` —— 审核通过商品 -> 千问上架文案。

纯独立函数，**不包含 agent 工具调用逻辑**：

* 入口：``product_select_result`` 中 ``audit_status='pass'`` 的商品
  （可选 ``product_ids`` 二次圈定）；hold/reject 商品**不注入**；
* 事实卡：按 ``product_id`` 关联 ``product_info``（标题/类目/价格/销量/
  颜色SKU/店铺背书），字段缺失的商品直接记 ``failed`` 不阻断批次；
* 模型：百炼千问 ``qwen3.7-flash-2026-07-15``（MultiModalConversation，
  与 ai_selector 同调用方式）；密钥只读 ``DASHSCOPE_API_KEY`` 环境变量；
* 三层校验：① JSON 结构 ② 硬约束（字数/条数/禁用词/候选互异）
  ③ 事实核对（specs 必须是事实卡子串、标题数字必须来自事实卡）；
  校验失败带原因重试（``MAX_RETRY`` 次）；
* 双状态维度（表 ``product_material``，一行一商品）：
  - ``gen_status``：pending -> generating -> success/failed（**断点续跑依据**，
    重跑幂等：跳过 success、补 failed/未生成；重新 success 时 audit 重置待审）；
  - ``audit_status``：NULL 待审 -> pass/hold/reject（人工填写，经
    ``excel_material_import`` 回写，状态机见 :mod:`src.material.material_audit`）；
* 输出：MySQL ``product_material`` + ``output/{ts}_material_gen.json``（调试）
  + ``output/excel/{ts}_pending_material_audit.xlsx``（待人工审核）。

用法::

    from src.material.generator import generate_goods_material
    generate_goods_material()                    # 全部 pass 商品（幂等续跑）
    generate_goods_material(product_ids=[1, 2])  # 只处理指定商品
"""
from __future__ import annotations

import json
import os
import re
import sys
import time
from datetime import datetime
from http import HTTPStatus
from pathlib import Path
from typing import Any

from dashscope import MultiModalConversation
from dashscope.common.error import DashScopeException, TimeoutException
from openpyxl import Workbook

# sys.path 引导：支持 `python src/material/generator.py` 直接启动
_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.common import database  # noqa: E402
from src.settings import OUTPUT_DIR  # noqa: E402

# ---- 模型与提示词版本（免费额度有限，锁定日期版本便于对账） ----
MODEL = "qwen3.7-flash-2026-07-15"
# 新版百炼业务空间 ID：只从环境变量读取（config/.env），禁止硬编码；
# 留空 = 使用账号默认业务空间（不携带 workspace 参数）
WORKSPACE_ID = os.environ.get("DASHSCOPE_WORKSPACE_ID", "").strip()
PROMPT_VERSION = "PV-20261011-1"

# 校验失败重试次数（1 次生成 + MAX_RETRY 次修正重试）
MAX_RETRY = 2
# 商品间间隔（免费档限流保护）
ITEM_DELAY_SECONDS = 0.5

TABLE = "product_material"
SELECT_TABLE = "product_select_result"
INFO_TABLE = "product_info"

# ---- 文案硬约束 ----
TITLE_MAX = 30
TITLE_COUNT = 3
POINT_MAX = 20
POINT_COUNT = 5

# 广告法绝对化用语及极限词（标题/卖点/详情骨架零容忍）
BANNED_WORDS = (
    "最", "最佳", "最优", "最好", "第一", "No.1", "no.1", "TOP1", "Top1", "top1",
    "顶级", "国家级", "世界级", "永久", "100%", "绝对", "独家",
    "全网最低", "史上最强", "疯狂", "无敌", "神器",
)

# ---- 待审核 Excel 列（末尾三列留空交人工填写；与选品审核同名约定） ----
EXCEL_COLUMNS = (
    "id", "product_id", "source_title", "cate_name", "wholesale_price",
    "title_v1", "title_v2", "title_v3",
    "point_1", "point_2", "point_3", "point_4", "point_5",
    "detail_pain", "detail_scene", "detail_specs", "detail_trust",
    "brief_scene", "brief_style", "brief_palette", "brief_point",
    "machine_check", "risk_flags",
    "audit_status", "audit_operator", "audit_comment",
)

_CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS `product_material` (
    `id`                BIGINT          NOT NULL AUTO_INCREMENT,
    `product_id`        BIGINT          NOT NULL                COMMENT '商品ID(product_info.product_id)',
    `gen_status`        VARCHAR(16)     NOT NULL DEFAULT 'pending' COMMENT '机器生成状态: pending/generating/success/failed',
    `gen_error`         VARCHAR(500)    NULL                    COMMENT '生成失败原因(gen_status=failed)',
    `content_json`      JSON            NULL                    COMMENT '千问生成的文案JSON',
    `model`             VARCHAR(64)     NULL                    COMMENT '生成所用模型',
    `prompt_version`    VARCHAR(32)     NULL                    COMMENT '提示词版本(调优钩子)',
    `machine_check`     VARCHAR(500)    NULL                    COMMENT '机器校验结论: pass 或 fail:原因摘要',
    `risk_flags_json`   JSON            NULL                    COMMENT '模型自报存疑点JSON数组',
    `audit_status`      VARCHAR(16)     NULL                    COMMENT '人工审核状态: NULL=待审 pass/hold/reject',
    `audit_operator`    VARCHAR(64)     NULL                    COMMENT '人工审核人',
    `audit_comment`     TEXT            NULL                    COMMENT '人工审核备注',
    `audit_time`        DATETIME        NULL                    COMMENT '人工审核时间',
    `created_at`        DATETIME        NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '写入时间',
    `updated_at`        DATETIME        NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '最近更新时间',
    PRIMARY KEY (`id`),
    UNIQUE KEY `uk_product_id` (`product_id`),
    KEY `idx_gen_status` (`gen_status`),
    KEY `idx_audit_status` (`audit_status`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_general_ci
  COMMENT='AI素材(文案)生成表: 双状态维度(gen_status机器/audit_status人工)'
"""

# ---- 内置 system 系统提示词（结构锁死 + 铁律防编造/违禁） ----
_SYSTEM_PROMPT = (
    "你是资深电商上架文案专家，为「货源商品 -> 自有店铺上架」场景撰写标题与卖点。\n"
    "\n"
    "【铁律，违反任何一条即为无效输出】\n"
    "1. 只允许使用【商品事实】中给出的信息，禁止编造功效、材质、认证、销量、"
    "适用人群等任何未给出的信息；不确定宁可不写。\n"
    f"2. 禁用词（广告法绝对化用语及极限词，标题/卖点/详情骨架中零容忍）：{'、'.join(BANNED_WORDS)}\n"
    f"3. 字数硬约束：主标题不超过{TITLE_MAX}个字符；每条短卖点不超过{POINT_MAX}个字符；"
    f"短卖点恰好{POINT_COUNT}条；主标题恰好{TITLE_COUNT}个候选。\n"
    "4. 标题结构=类目词+属性词+场景词+卖点词；不出现【】#等堆砌符号；"
    "3个标题策略必须各异（候选1搜索流量向、候选2场景转化向、候选3卖点差异向）。\n"
    "5. detail_outline.specs 只允许从商品事实的参数行**原样摘录**（如颜色/尺码/价格/评分），"
    "禁止新增或改写。\n"
    "6. brief 是给后续 AI 生图用的场景简报（不面向消费者，不受禁用词约束）："
    "scene 描述一个具体使用场景画面，palette 给配色建议，selling_point 选一个"
    "最值得视觉化的卖点。\n"
    "7. 输出只允许是一个 JSON 对象，不要 markdown 代码块，不要任何解释文字。\n"
    "\n"
    "【输出 JSON 规范（字段必须齐全）】\n"
    "{\n"
    '  "product_id": <整数，原样回传>,\n'
    f'  "titles": [{{"version": 1, "title": "<={TITLE_MAX}字>", "strategy": "搜索"}}, '
    f'{{"version": 2, "title": "...", "strategy": "场景"}}, '
    f'{{"version": 3, "title": "...", "strategy": "卖点"}}],\n'
    f'  "short_points": ["<={POINT_MAX}字>", "<={POINT_MAX}字>", "<={POINT_MAX}字>", '
    f'"<={POINT_MAX}字>", "<={POINT_MAX}字>"],\n'
    '  "detail_outline": {"pain_point": "<痛点1句>", "scene": "<场景1句>", '
    '"specs": ["<事实卡参数原样摘录>"], "trust": "<信任背书1句，可空字符串>"},\n'
    '  "sku_renames": [{"sku": "<原SKU名>", "new_name": "<营销化改名>"}],\n'
    '  "brief": {"scene": "<生图场景画面>", "style": "<风格调性>", '
    '"palette": "<配色建议>", "selling_point": "<最值得视觉化的卖点>"},\n'
    '  "risk_flags": ["<对事实卡存疑、建议人工复核的点；无则空数组>"]\n'
    "}\n"
    "\n"
    "【示例】\n"
    '{"product_id": 4643797, "titles": '
    '[{"version": 1, "title": "男士透气网面运动鞋轻便防滑跑步鞋春季新款", "strategy": "搜索"}, '
    '{"version": 2, "title": "春季晨跑必备男士网面透气运动鞋轻便舒适", "strategy": "场景"}, '
    '{"version": 3, "title": "EVA缓震男士运动鞋橡胶防滑大底跑步鞋", "strategy": "卖点"}], '
    '"short_points": ["网面透气鞋面干爽不闷脚", "EVA缓震中底回弹轻盈", '
    '"橡胶防滑大底抓地稳", "轻量设计单只约260g", "多色多码通勤慢跑都好搭"], '
    '"detail_outline": {"pain_point": "长时间行走脚部闷热易疲劳", '
    '"scene": "适合春季晨跑与日常通勤穿着", '
    '"specs": ["黑色、灰色、藏青", "尺码：38-44", "拿货价：48 元"], '
    '"trust": "源头工厂直供，支持一件代发"}, '
    '"sku_renames": [{"sku": "黑色", "new_name": "经典黑"}], '
    '"brief": {"scene": "清晨城市公园跑道，运动鞋特写", "style": "活力运动", '
    '"palette": "浅灰+天空蓝", "selling_point": "橡胶防滑大底纹路特写"}, '
    '"risk_flags": []}'
)


# --------------------------------------------------------------------------
# 纯函数：解析 / 事实卡 / 三层校验
# --------------------------------------------------------------------------
def _extract_text(resp) -> str:
    """从多模态响应中抽取纯文本（与 ai_selector 同构）。"""
    content = resp.output.choices[0].message.content
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = []
        for part in content:
            if isinstance(part, dict):
                parts.append(part.get("text") or "")
            else:
                parts.append(getattr(part, "text", "") or str(part))
        return "".join(parts)
    return str(content or "")


def _parse_json_object(text: str) -> dict:
    """剥掉代码围栏等包装，解析出 JSON 对象（失败抛异常）。"""
    cleaned = text.strip()
    fence = re.match(r"^```(?:json)?\s*(.*?)\s*```$", cleaned, re.S)
    if fence:
        cleaned = fence.group(1).strip()
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start < 0 or end <= start:
            raise
        data = json.loads(cleaned[start:end + 1])
    if not isinstance(data, dict):
        raise ValueError("模型输出不是 JSON 对象")
    return data


def _find_banned(text: str) -> str | None:
    """返回文本中命中的第一个禁用词（长词优先，反馈更精确；无命中返回 None）。"""
    for word in sorted(BANNED_WORDS, key=len, reverse=True):
        if word in text:
            return word
    return None


def _digit_runs(text: str) -> list[str]:
    """提取长度 >= 2 的连续数字串（事实核对用）。"""
    return re.findall(r"\d{2,}", text)


def _loads(value: Any) -> Any:
    """JSON 字符串 -> 对象（DB JSON 列以 str 返回）；已是对象则原样返回。"""
    if isinstance(value, (list, dict)):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return None
    return None


def _assemble_fact_card(row: dict) -> str:
    """把 product_info 一行组装成给千问看的事实卡（字段缺失行省略）。

    :raises ValueError: 标题为空（该商品直接判 failed 的依据）。
    """
    name = str(row.get("product_name") or "").strip()
    if not name:
        raise ValueError("product_info.product_name 为空")
    lines = [
        f"- 商品ID：{row.get('product_id')}",
        f"- 商品标题：{name}",
    ]
    cate = str(row.get("cate_name") or "").strip()
    if cate:
        lines.append(f"- 类目：{cate}")
    for label, key, suffix in (
        ("批发价", "wholesale_price", " 元"),
        ("零售价", "retail_price", " 元"),
    ):
        val = row.get(key)
        if val is not None:
            lines.append(f"- {label}：{val}{suffix}")
    buy = row.get("buy_count")
    if buy is not None:
        lines.append(f"- 销量：{buy}")
    stock_total = row.get("stock_total")
    if stock_total is not None:
        lines.append(f"- 可售库存合计：{stock_total}")
    ret = row.get("return_goods_rate")
    if ret is not None:
        try:
            lines.append(f"- 退货率：{round(float(ret) * 100, 2)}%")
        except (TypeError, ValueError):
            pass
    colors = _loads(row.get("colors"))
    color_names = [
        str(c.get("name")).strip()
        for c in colors
        if isinstance(c, dict) and c.get("name")
    ] if isinstance(colors, list) else []
    if color_names:
        lines.append(f"- 颜色分类：{'、'.join(color_names[:8])}")
    sku_stock = _loads(row.get("sku_stock"))
    if isinstance(sku_stock, list) and sku_stock:
        sku_bits = []
        for s in sku_stock[:4]:
            if not isinstance(s, dict):
                continue
            bit = str(s.get("name") or "").strip()
            price = s.get("price")
            stock = s.get("stock")
            detail = " ".join(
                x for x in (
                    f"价{price}" if price is not None else "",
                    f"库存{stock}" if stock is not None else "",
                ) if x
            )
            sku_bits.append(f"{bit}({detail})" if detail else bit)
        if sku_bits:
            lines.append(f"- SKU示例：{'；'.join(sku_bits)}")
    shop = str(row.get("shop_name") or "").strip()
    market = str(row.get("market_name") or "").strip()
    area = str(row.get("market_area_name") or "").strip()
    if shop or market:
        loc = market + (f" {area}" if area else "")
        lines.append(f"- 店铺/档口：{shop or '-'}（{loc}）" if loc else f"- 店铺/档口：{shop}")
    score = row.get("rate_trade_score")
    if score is not None:
        lines.append(f"- 店铺交易评分：{score}（5分制）")
    h24 = row.get("h24_delivery_rate")
    if h24 is not None:
        try:
            lines.append(f"- 24小时发货率：{round(float(h24) * 100, 2)}%")
        except (TypeError, ValueError):
            pass
    shop_age = row.get("shop_age")
    if shop_age is not None:
        lines.append(f"- 店龄：{shop_age} 年")
    lines.append("- 货源背书：源头档口直供，支持一件代发")
    return "\n".join(lines)


def _validate_material(data: dict, expected_pid: int, fact_card: str) -> dict:
    """三层校验：结构 -> 硬约束（字数/条数/禁用词）-> 事实核对。

    :returns: 归一化后的文案内容字典（落库 content_json 的结构）。
    :raises ValueError: 任一层不通过（message 带定位信息，供重试反馈）。
    """
    # ---- ① 结构层 ----
    if not isinstance(data, dict):
        raise ValueError("输出不是 JSON 对象")
    try:
        pid = int(data.get("product_id"))
    except (TypeError, ValueError) as exc:
        raise ValueError("product_id 缺失或非整数") from exc
    if pid != int(expected_pid):
        raise ValueError(f"product_id 不匹配：期望 {expected_pid}，实际 {pid}")

    titles_raw = data.get("titles")
    if not isinstance(titles_raw, list) or len(titles_raw) != TITLE_COUNT:
        raise ValueError(f"titles 必须是恰好 {TITLE_COUNT} 个元素的数组")
    points_raw = data.get("short_points")
    if not isinstance(points_raw, list) or len(points_raw) != POINT_COUNT:
        raise ValueError(f"short_points 必须是恰好 {POINT_COUNT} 个元素的数组")
    outline_raw = data.get("detail_outline")
    if not isinstance(outline_raw, dict):
        raise ValueError("detail_outline 必须是对象")
    brief_raw = data.get("brief")
    if not isinstance(brief_raw, dict):
        raise ValueError("brief 必须是对象")
    skus_raw = data.get("sku_renames")
    if skus_raw is None:
        skus_raw = []
    if not isinstance(skus_raw, list):
        raise ValueError("sku_renames 必须是数组")
    flags_raw = data.get("risk_flags")
    if flags_raw is None:
        flags_raw = []
    if not isinstance(flags_raw, list):
        raise ValueError("risk_flags 必须是数组")

    # ---- ② 硬约束层：标题 ----
    fact_digits = set(_digit_runs(fact_card))
    titles: list[dict] = []
    seen_titles: set[str] = set()
    for idx, item in enumerate(titles_raw, 1):
        if not isinstance(item, dict):
            raise ValueError(f"titles[{idx}] 不是对象")
        title = str(item.get("title") or "").strip()
        if not title:
            raise ValueError(f"titles[{idx}].title 为空")
        if len(title) > TITLE_MAX:
            raise ValueError(f"titles[{idx}] 超长（{len(title)}>{TITLE_MAX}）: {title!r}")
        banned = _find_banned(title)
        if banned:
            raise ValueError(f"titles[{idx}] 命中禁用词 {banned!r}: {title!r}")
        if title in seen_titles:
            raise ValueError(f"titles[{idx}] 与前面候选重复: {title!r}")
        seen_titles.add(title)
        for digit in _digit_runs(title):
            if digit not in fact_digits:
                raise ValueError(
                    f"titles[{idx}] 含事实卡中不存在的数字 {digit!r}（疑似编造）: {title!r}"
                )
        try:
            version = int(item.get("version") or idx)
        except (TypeError, ValueError):
            version = idx
        titles.append({
            "version": version,
            "title": title,
            "strategy": str(item.get("strategy") or "").strip(),
        })

    # ---- ② 硬约束层：短卖点 ----
    points: list[str] = []
    for idx, raw in enumerate(points_raw, 1):
        point = str(raw or "").strip()
        if not point:
            raise ValueError(f"short_points[{idx}] 为空")
        if len(point) > POINT_MAX:
            raise ValueError(
                f"short_points[{idx}] 超长（{len(point)}>{POINT_MAX}）: {point!r}"
            )
        banned = _find_banned(point)
        if banned:
            raise ValueError(f"short_points[{idx}] 命中禁用词 {banned!r}: {point!r}")
        points.append(point)

    # ---- ③ 事实核对层：详情骨架 ----
    pain = str(outline_raw.get("pain_point") or "").strip()
    if not pain:
        raise ValueError("detail_outline.pain_point 为空")
    banned = _find_banned(pain)
    if banned:
        raise ValueError(f"pain_point 命中禁用词 {banned!r}")
    scene = str(outline_raw.get("scene") or "").strip()
    if not scene:
        raise ValueError("detail_outline.scene 为空")
    banned = _find_banned(scene)
    if banned:
        raise ValueError(f"scene 命中禁用词 {banned!r}")
    trust = str(outline_raw.get("trust") or "").strip()
    banned = _find_banned(trust)
    if banned:
        raise ValueError(f"trust 命中禁用词 {banned!r}")
    specs_raw = outline_raw.get("specs")
    if specs_raw is None:
        specs_raw = []
    if not isinstance(specs_raw, list):
        raise ValueError("detail_outline.specs 必须是数组")
    specs: list[str] = []
    for idx, raw in enumerate(specs_raw, 1):
        spec = str(raw or "").strip()
        if not spec:
            continue
        if spec not in fact_card:
            raise ValueError(
                f"specs[{idx}] 不是事实卡原样摘录（疑似编造/改写）: {spec!r}"
            )
        specs.append(spec)

    # ---- SKU 改名（可空；改名保留原名识别信息由人工终审把关） ----
    sku_renames: list[dict] = []
    for raw in skus_raw:
        if not isinstance(raw, dict):
            continue
        sku = str(raw.get("sku") or "").strip()
        new_name = str(raw.get("new_name") or "").strip()
        if sku and new_name:
            sku_renames.append({"sku": sku, "new_name": new_name})

    # ---- brief（生图简报；不面向消费者，不做禁用词校验） ----
    brief = {}
    for key in ("scene", "style", "palette", "selling_point"):
        val = str(brief_raw.get(key) or "").strip()
        if not val:
            raise ValueError(f"brief.{key} 为空")
        brief[key] = val

    risk_flags = [str(f).strip() for f in flags_raw if str(f).strip()]

    return {
        "product_id": pid,
        "titles": titles,
        "short_points": points,
        "detail_outline": {
            "pain_point": pain,
            "scene": scene,
            "specs": specs,
            "trust": trust,
        },
        "sku_renames": sku_renames,
        "brief": brief,
        "risk_flags": risk_flags,
    }


# --------------------------------------------------------------------------
# 模型调用（单商品；重试循环在 _generate_one）
# --------------------------------------------------------------------------
def _call_qwen(fact_card: str, feedback: str = "") -> str:
    """调用千问生成文案，返回原始文本；网络/接口异常抛 RuntimeError。"""
    user_text = (
        "【商品事实】\n"
        f"{fact_card}\n\n"
        "请按 System 规范输出 JSON 对象，只输出 JSON 本身，"
        "不要 markdown 代码块，不要任何解释文字。"
    )
    if feedback:
        user_text += (
            f"\n\n【修正要求】你上一次的输出未通过校验：{feedback}。"
            "请修正后重新输出完整 JSON。"
        )
    messages = [
        {"role": "system", "content": [{"text": _SYSTEM_PROMPT}]},
        {"role": "user", "content": [{"text": user_text}]},
    ]
    try:
        call_kwargs: dict = {
            "model": MODEL,
            "messages": messages,
            "enable_thinking": False,
        }
        if WORKSPACE_ID:  # 配置了业务空间才携带 workspace 参数
            call_kwargs["workspace"] = WORKSPACE_ID
        resp = MultiModalConversation.call(**call_kwargs)
    except (TimeoutError, ConnectionError, OSError) as exc:
        raise RuntimeError(f"网络异常 {type(exc).__name__}: {exc}") from exc
    except (TimeoutException, DashScopeException) as exc:
        raise RuntimeError(f"模型调用异常 {type(exc).__name__}: {exc}") from exc
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"模型调用异常 {type(exc).__name__}: {exc}") from exc

    if getattr(resp, "status_code", None) != HTTPStatus.OK:
        raise RuntimeError(
            f"模型调用失败 status_code={getattr(resp, 'status_code', None)} "
            f"message={getattr(resp, 'message', resp)}"
        )
    try:
        raw_text = _extract_text(resp)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"响应结构解析失败 {type(exc).__name__}: {exc}") from exc
    if not raw_text.strip():
        raise RuntimeError("模型返回内容为空")
    return raw_text


def _generate_one(pid: int, info_row: dict) -> tuple[bool, dict | str]:
    """单商品：事实卡 -> 千问 -> 三层校验（失败带原因重试）。

    :returns: ``(True, 归一化文案内容)`` 或 ``(False, 失败原因)``。
    """
    try:
        fact_card = _assemble_fact_card(info_row)
    except Exception as exc:  # noqa: BLE001 - 事实卡组装失败即判 failed
        return False, f"事实卡组装失败: {exc}"

    feedback = ""
    last_err = "未执行"
    for _attempt in range(1 + MAX_RETRY):
        try:
            raw = _call_qwen(fact_card, feedback)
        except RuntimeError as exc:
            last_err = str(exc)
            feedback = last_err
            continue
        try:
            data = _parse_json_object(raw)
            content = _validate_material(data, pid, fact_card)
            return True, content
        except (json.JSONDecodeError, ValueError) as exc:
            last_err = str(exc)
            feedback = last_err
    return False, last_err


# --------------------------------------------------------------------------
# DB 生命周期（gen_status 断点）
# --------------------------------------------------------------------------
_table_ready = False


def _ensure_table() -> None:
    """幂等建表（进程内只执行一次 DDL）。"""
    global _table_ready
    if _table_ready:
        return
    with database.connection() as conn:
        with conn.cursor() as cur:
            cur.execute(_CREATE_TABLE_SQL)
    _table_ready = True


def _query_pass_product_ids(product_ids: list | None = None) -> list[int]:
    """查选品审核 pass 的去重商品 ID（可选与入参取交集）。"""
    rows = database.fetch_all(
        SELECT_TABLE,
        where="audit_status=%s",
        params=["pass"],
        cols="product_id",
    )
    pass_ids = sorted({int(r["product_id"]) for r in rows if r.get("product_id")})
    if not product_ids:
        return pass_ids
    wanted: set[int] = set()
    for x in product_ids:
        try:
            wanted.add(int(x))
        except (TypeError, ValueError):
            continue
    return [pid for pid in pass_ids if pid in wanted]


def _query_info_rows(pids: list[int]) -> dict[int, dict]:
    """按商品 ID 查 product_info（同商品多行取 created_at 最新一行）。"""
    if not pids:
        return {}
    placeholders = ", ".join(["%s"] * len(pids))
    rows = database.fetch_all(
        INFO_TABLE,
        where=f"product_id IN ({placeholders})",
        params=pids,
    )
    best: dict[int, dict] = {}
    for row in rows:
        try:
            pid = int(row.get("product_id"))
        except (TypeError, ValueError):
            continue
        prev = best.get(pid)
        if prev is None or str(row.get("created_at") or "") > str(prev.get("created_at") or ""):
            best[pid] = row
    return best


def _query_material_rows(pids: list[int]) -> dict[int, dict]:
    """按商品 ID 查 product_material 现有行。"""
    if not pids:
        return {}
    placeholders = ", ".join(["%s"] * len(pids))
    rows = database.fetch_all(
        TABLE,
        where=f"product_id IN ({placeholders})",
        params=pids,
    )
    out: dict[int, dict] = {}
    for row in rows:
        try:
            out[int(row.get("product_id"))] = row
        except (TypeError, ValueError):
            continue
    return out


def _insert_pending_rows(pids: list[int]) -> int:
    """预写 pending 行（仅补不存在的商品，不动已有行）。"""
    rows = [{"product_id": pid, "gen_status": "pending"} for pid in pids]
    return database.insert_many(TABLE, rows) if rows else 0


def _mark_generating(pid: int) -> None:
    with database.connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE `product_material` SET `gen_status`='generating', "
                "`gen_error`=NULL WHERE `product_id`=%s",
                (pid,),
            )


def _mark_success(pid: int, content: dict, machine_check: str) -> None:
    """落库成功行；audit 四字段重置（新版本素材需重新人工审核）。"""
    with database.connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE `product_material` SET `gen_status`='success', `gen_error`=NULL, "
                "`content_json`=%s, `model`=%s, `prompt_version`=%s, `machine_check`=%s, "
                "`risk_flags_json`=%s, `audit_status`=NULL, `audit_operator`=NULL, "
                "`audit_comment`=NULL, `audit_time`=NULL "
                "WHERE `product_id`=%s",
                (
                    json.dumps(content, ensure_ascii=False),
                    MODEL,
                    PROMPT_VERSION,
                    machine_check,
                    json.dumps(content.get("risk_flags") or [], ensure_ascii=False),
                    pid,
                ),
            )


def _mark_failed(pid: int, error: str) -> None:
    with database.connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE `product_material` SET `gen_status`='failed', `gen_error`=%s "
                "WHERE `product_id`=%s",
                (str(error)[:490], pid,),
            )


# --------------------------------------------------------------------------
# 产物：待审核 Excel + 调试 JSON
# --------------------------------------------------------------------------
def _content_to_columns(content: dict) -> dict:
    """归一化文案内容 -> Excel 行片段（title_v1..v3 / point_1..5 / detail / brief）。"""
    titles = content.get("titles") or []
    points = content.get("short_points") or []
    outline = content.get("detail_outline") or {}
    brief = content.get("brief") or {}
    row = {}
    for i in range(TITLE_COUNT):
        row[f"title_v{i + 1}"] = titles[i].get("title") if i < len(titles) else None
    for i in range(POINT_COUNT):
        row[f"point_{i + 1}"] = points[i] if i < len(points) else None
    row["detail_pain"] = outline.get("pain_point")
    row["detail_scene"] = outline.get("scene")
    row["detail_specs"] = "；".join(outline.get("specs") or [])
    row["detail_trust"] = outline.get("trust")
    row["brief_scene"] = brief.get("scene")
    row["brief_style"] = brief.get("style")
    row["brief_palette"] = brief.get("palette")
    row["brief_point"] = brief.get("selling_point")
    return row


def _export_audit_excel(
    material_rows: list[dict], info_rows: dict[int, dict], ts: str,
) -> Path:
    """导出待人工审核 Excel（首列 id + 内容列 + 末尾三列留空）。"""
    excel_dir = OUTPUT_DIR / "excel"
    excel_dir.mkdir(parents=True, exist_ok=True)
    path = excel_dir / f"{ts}_pending_material_audit.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "素材待审核"
    ws.append(list(EXCEL_COLUMNS))
    for m in sorted(material_rows, key=lambda r: int(r.get("id") or 0)):
        content = _loads(m.get("content_json")) or {}
        if not isinstance(content, dict):
            continue
        info = info_rows.get(int(m.get("product_id") or 0)) or {}
        row = {
            "id": m.get("id"),
            "product_id": m.get("product_id"),
            "source_title": info.get("product_name"),
            "cate_name": info.get("cate_name"),
            "wholesale_price": info.get("wholesale_price"),
            **_content_to_columns(content),
            "machine_check": m.get("machine_check"),
            "risk_flags": "；".join(_loads(m.get("risk_flags_json")) or [])
            if isinstance(_loads(m.get("risk_flags_json")), list)
            else None,
            "audit_status": None,
            "audit_operator": None,
            "audit_comment": None,
        }
        ws.append([row.get(col) for col in EXCEL_COLUMNS])
    for idx, col in enumerate(EXCEL_COLUMNS, 1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = (
            18 if col.startswith("brief_") or col.startswith("title_")
            else 14 if col.startswith("audit_") or col.startswith("point_")
            else max(12, min(40, len(col) + 4))
        )
    wb.save(path)
    return path


def _dump_debug_json(payload: dict, ts: str) -> Path:
    """调试 JSON：整批生成明细（模型原始输出不在内，内容已归一化）。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"{ts}_material_gen.json"
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    return path


# --------------------------------------------------------------------------
# 主入口
# --------------------------------------------------------------------------
def generate_goods_material(
    product_ids: list | None = None,
    export_excel: bool = True,
) -> dict:
    """P1 文案生成主流程：pass 商品 -> 事实卡 -> 千问 -> 校验 -> 落库/Excel。

    :param product_ids:  可选二次圈定（与 DB 中 pass 集合取交集）；
                         None/空 = 全部 ``audit_status='pass'`` 商品。
    :param export_excel: False 跳过 Excel/JSON 落盘（单测用）。
    :returns: ``{"success", "count", "generated", "failed", "skipped_success",
                "excel_path", "json_path", "failures", "message"}``
    """
    _ensure_table()
    pids = _query_pass_product_ids(product_ids)
    if not pids:
        return {
            "success": False,
            "count": 0,
            "generated": 0,
            "failed": 0,
            "skipped_success": 0,
            "excel_path": None,
            "json_path": None,
            "failures": [],
            "message": "无审核通过（audit_status='pass'）商品可生成素材",
        }

    info_rows = _query_info_rows(pids)
    material_rows = _query_material_rows(pids)

    # ---- 预写 pending（仅补缺行） ----
    missing = [pid for pid in pids if pid not in material_rows]
    if missing:
        _insert_pending_rows(missing)
        for pid in missing:
            material_rows[pid] = {"product_id": pid, "gen_status": "pending"}

    # ---- 断点过滤：跳过 success（幂等续跑），重跑 failed/卡死 generating/未生成 ----
    todo = [
        pid for pid in pids
        if str(material_rows.get(pid, {}).get("gen_status") or "pending") != "success"
    ]
    skipped = len(pids) - len(todo)
    if skipped:
        print(f"[素材生成] 断点命中：跳过已 success {skipped} 个，待生成 {len(todo)} 个")

    generated = 0
    failed = 0
    failures: list[dict] = []
    for seq, pid in enumerate(todo, 1):
        info_row = info_rows.get(pid)
        print(f"[素材生成] ({seq}/{len(todo)}) product_id={pid} ...")
        _mark_generating(pid)
        if not info_row or not str(info_row.get("product_name") or "").strip():
            reason = "product_info 缺失或标题为空"
            _mark_failed(pid, reason)
            failed += 1
            failures.append({"product_id": pid, "error": reason})
            print(f"[素材生成]   -> failed: {reason}")
            continue
        ok, result = _generate_one(pid, info_row)
        if ok:
            _mark_success(pid, result, "pass")
            generated += 1
            print(f"[素材生成]   -> success（标题候选 {TITLE_COUNT} / 卖点 {POINT_COUNT}）")
        else:
            _mark_failed(pid, result)
            failed += 1
            failures.append({"product_id": pid, "error": result})
            print(f"[素材生成]   -> failed: {result}")
        if seq < len(todo) and ITEM_DELAY_SECONDS > 0:
            time.sleep(ITEM_DELAY_SECONDS)

    # ---- 审核产物：重查 success 且待审的行 ----
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    refreshed = _query_material_rows(pids)
    auditable = [
        row for pid, row in refreshed.items()
        if str(row.get("gen_status")) == "success" and not row.get("audit_status")
    ]
    excel_path: Path | None = None
    json_path: Path | None = None
    if export_excel and auditable:
        excel_path = _export_audit_excel(auditable, info_rows, ts)
        print(f"[素材生成] 待审核 Excel: output/excel/{excel_path.name}（{len(auditable)} 行）")
    if export_excel:
        json_path = _dump_debug_json({
            "generated_at": ts,
            "model": MODEL,
            "prompt_version": PROMPT_VERSION,
            "target_count": len(pids),
            "todo_count": len(todo),
            "generated": generated,
            "failed": failed,
            "skipped_success": skipped,
            "failures": failures,
        }, ts)
        print(f"[素材生成] 调试 JSON: output/{json_path.name}")

    success = generated > 0 or (not todo and skipped > 0)
    message = (
        f"文案生成完成：目标 {len(pids)} 个 / 本次生成 {generated} / "
        f"断点跳过 {skipped} / 失败 {failed}"
        + (f"；待审核 Excel: output/excel/{excel_path.name}" if excel_path else "")
    )
    print(f"[素材生成] {message}")
    return {
        "success": success,
        "count": len(pids),
        "generated": generated,
        "failed": failed,
        "skipped_success": skipped,
        "excel_path": str(excel_path) if excel_path else None,
        "json_path": str(json_path) if json_path else None,
        "failures": failures,
        "message": message,
    }


def main() -> int:
    """CLI：``python src/material/generator.py [--product-ids 1,2] [--no-excel]``。"""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    import argparse

    parser = argparse.ArgumentParser(description="P1 素材(文案)生成：审核通过商品 -> 千问文案")
    parser.add_argument(
        "--product-ids", default="",
        help="逗号分隔商品ID（可选，默认全部 audit_status='pass' 商品）",
    )
    parser.add_argument(
        "--no-excel", action="store_true",
        help="只生成入库，不导出待审核 Excel/JSON",
    )
    args = parser.parse_args()
    pids = [int(x) for x in re.split(r"[,\s]+", args.product_ids) if x.strip()] or None
    result = generate_goods_material(product_ids=pids, export_excel=not args.no_excel)
    return 0 if result.get("success") else 1


if __name__ == "__main__":
    sys.exit(main())
