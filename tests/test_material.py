# -*- coding: utf-8 -*-
"""P1 素材(文案)生成单元测试：纯函数校验 + 状态机 + 导出，全程无网络无真实 DB。"""
from __future__ import annotations

import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

from openpyxl import Workbook, load_workbook

from src.common import database
from src.material import excel_material_import as emi
from src.material import generator
from src.material.excel_material_import import import_material_audit_excel
from src.material.generator import (
    BANNED_WORDS,
    EXCEL_COLUMNS,
    POINT_COUNT,
    TITLE_COUNT,
    _assemble_fact_card,
    _content_to_columns,
    _digit_runs,
    _export_audit_excel,
    _find_banned,
    _generate_one,
    _parse_json_object,
    _validate_material,
    generate_goods_material,
)
from src.material.material_audit import material_audit_product


def _full_row() -> dict:
    """字段齐全的 product_info 行样例。"""
    return {
        "product_id": 4643797,
        "product_name": "2026春季新款男士透气网面运动鞋轻便防滑跑步鞋",
        "cate_name": "男鞋/运动鞋/跑鞋",
        "wholesale_price": 48,
        "retail_price": 129,
        "buy_count": 1200,
        "stock_total": 3560,
        "return_goods_rate": 0.032,
        "colors": json.dumps([{"vid": 1, "name": "黑色"}, {"vid": 2, "name": "灰色"}]),
        "sku_stock": json.dumps([
            {"name": "黑色/41", "price": 48, "stock": 120},
            {"name": "灰色/42", "price": 48, "stock": 80},
        ]),
        "shop_name": "某某鞋业档口",
        "market_name": "国际鞋都",
        "market_area_name": "A区3楼",
        "rate_trade_score": 4.8,
        "h24_delivery_rate": 0.99,
        "shop_age": 6,
    }


def _good_material(pid: int = 4643797) -> dict:
    """结构合法的模型输出（与 _SYSTEM_PROMPT 示例同构）。"""
    return {
        "product_id": pid,
        "titles": [
            {"version": 1, "title": "男士透气网面运动鞋轻便防滑跑步鞋春季新款", "strategy": "搜索"},
            {"version": 2, "title": "春季晨跑必备男士网面透气运动鞋轻便舒适", "strategy": "场景"},
            {"version": 3, "title": "EVA缓震男士运动鞋橡胶防滑大底跑步鞋", "strategy": "卖点"},
        ],
        "short_points": [
            "网面透气鞋面干爽不闷脚",
            "EVA缓震中底回弹轻盈",
            "橡胶防滑大底抓地稳",
            "轻量设计穿着无负担",
            "多色多码通勤慢跑都好搭",
        ],
        "detail_outline": {
            "pain_point": "长时间行走脚部闷热易疲劳",
            "scene": "适合春季晨跑与日常通勤穿着",
            "specs": ["批发价：48 元", "黑色、灰色"],
            "trust": "源头档口直供，支持一件代发",
        },
        "sku_renames": [{"sku": "黑色", "new_name": "经典黑"}],
        "brief": {
            "scene": "清晨城市公园跑道，运动鞋特写",
            "style": "活力运动",
            "palette": "浅灰+天空蓝",
            "selling_point": "橡胶防滑大底纹路特写",
        },
        "risk_flags": [],
    }


class ParseJsonTests(unittest.TestCase):
    def test_plain_object(self) -> None:
        self.assertEqual(_parse_json_object('{"a": 1}'), {"a": 1})

    def test_fenced_object(self) -> None:
        self.assertEqual(_parse_json_object('```json\n{"a": 1}\n```'), {"a": 1})

    def test_object_with_surrounding_text(self) -> None:
        data = _parse_json_object('好的，结果如下：{"a": 1} 请查收')
        self.assertEqual(data, {"a": 1})

    def test_non_object_raises(self) -> None:
        with self.assertRaises(ValueError):
            _parse_json_object("[1, 2]")

    def test_garbage_raises(self) -> None:
        with self.assertRaises(Exception):
            _parse_json_object("完全不是JSON")


class BannedAndDigitsTests(unittest.TestCase):
    def test_banned_word_detected(self) -> None:
        self.assertEqual(_find_banned("全网最低价运动鞋"), "全网最低")
        self.assertIsNone(_find_banned("春季男士运动鞋"))

    def test_banned_constant_covers_absolute_terms(self) -> None:
        for word in ("最", "第一", "100%", "独家"):
            self.assertIn(word, BANNED_WORDS)

    def test_digit_runs_length_ge_2(self) -> None:
        self.assertEqual(_digit_runs("尺码38-44约260g"), ["38", "44", "260"])


class FactCardTests(unittest.TestCase):
    def test_full_row(self) -> None:
        card = _assemble_fact_card(_full_row())
        self.assertIn("商品标题：2026春季新款", card)
        self.assertIn("批发价：48 元", card)
        self.assertIn("退货率：3.2%", card)
        self.assertIn("颜色分类：黑色、灰色", card)
        self.assertIn("黑色/41(价48 库存120)", card)
        self.assertIn("24小时发货率：99.0%", card)
        self.assertIn("源头档口直供", card)

    def test_minimal_row(self) -> None:
        card = _assemble_fact_card({"product_id": 1, "product_name": "跑鞋"})
        self.assertIn("商品ID：1", card)
        self.assertNotIn("批发价", card)  # 缺失字段整行省略

    def test_empty_name_raises(self) -> None:
        with self.assertRaises(ValueError):
            _assemble_fact_card({"product_id": 1, "product_name": "  "})


class ValidateMaterialTests(unittest.TestCase):
    def setUp(self) -> None:
        self.fact = _assemble_fact_card(_full_row())

    def test_happy_path(self) -> None:
        content = _validate_material(_good_material(), 4643797, self.fact)
        self.assertEqual(content["product_id"], 4643797)
        self.assertEqual(len(content["titles"]), TITLE_COUNT)
        self.assertEqual(len(content["short_points"]), POINT_COUNT)
        self.assertEqual(content["brief"]["style"], "活力运动")
        self.assertIn("批发价：48 元", content["detail_outline"]["specs"])

    def test_product_id_mismatch(self) -> None:
        with self.assertRaisesRegex(ValueError, "不匹配"):
            _validate_material(_good_material(999), 4643797, self.fact)

    def test_banned_word_in_title(self) -> None:
        data = _good_material()
        data["titles"][0]["title"] = "全网最低男士运动鞋春季透气款"
        with self.assertRaisesRegex(ValueError, "禁用词"):
            _validate_material(data, 4643797, self.fact)

    def test_title_too_long(self) -> None:
        data = _good_material()
        data["titles"][0]["title"] = "男" * 31
        with self.assertRaisesRegex(ValueError, "超长"):
            _validate_material(data, 4643797, self.fact)

    def test_wrong_title_count(self) -> None:
        data = _good_material()
        data["titles"] = data["titles"][:2]
        with self.assertRaisesRegex(ValueError, "titles"):
            _validate_material(data, 4643797, self.fact)

    def test_wrong_point_count(self) -> None:
        data = _good_material()
        data["short_points"] = data["short_points"][:4]
        with self.assertRaisesRegex(ValueError, "short_points"):
            _validate_material(data, 4643797, self.fact)

    def test_duplicate_titles(self) -> None:
        data = _good_material()
        data["titles"][1]["title"] = data["titles"][0]["title"]
        with self.assertRaisesRegex(ValueError, "重复"):
            _validate_material(data, 4643797, self.fact)

    def test_invented_digit_rejected(self) -> None:
        data = _good_material()
        data["titles"][0]["title"] = "月销10万男士透气网面运动鞋春季新款"
        with self.assertRaisesRegex(ValueError, "数字"):
            _validate_material(data, 4643797, self.fact)

    def test_spec_not_in_fact_card(self) -> None:
        data = _good_material()
        data["detail_outline"]["specs"] = ["德国进口橡胶大底"]
        with self.assertRaisesRegex(ValueError, "specs"):
            _validate_material(data, 4643797, self.fact)

    def test_missing_brief_field(self) -> None:
        data = _good_material()
        del data["brief"]["palette"]
        with self.assertRaisesRegex(ValueError, "brief.palette"):
            _validate_material(data, 4643797, self.fact)


class GenerateOneTests(unittest.TestCase):
    def test_model_call_failure_marks_failed(self) -> None:
        with mock.patch.object(
            generator, "_call_qwen", side_effect=RuntimeError("模型调用异常: mock")
        ):
            ok, err = _generate_one(1, _full_row())
        self.assertFalse(ok)
        self.assertIn("模型调用异常", err)

    def test_fact_card_failure_short_circuits(self) -> None:
        ok, err = _generate_one(1, {"product_id": 1, "product_name": ""})
        self.assertFalse(ok)
        self.assertIn("事实卡", err)

    def test_retry_with_feedback_then_success(self) -> None:
        bad = _good_material()
        bad["titles"][0]["title"] = "全网最低男士运动鞋春季透气款"
        good = _good_material()
        calls: list[str] = []

        def _fake_call(fact_card: str, feedback: str = "") -> str:
            calls.append(feedback)
            return json.dumps(bad if len(calls) == 1 else good, ensure_ascii=False)

        with mock.patch.object(generator, "_call_qwen", side_effect=_fake_call):
            ok, content = _generate_one(4643797, _full_row())
        self.assertTrue(ok)
        self.assertEqual(len(calls), 2)
        self.assertIn("禁用词", calls[1])  # 第二次调用带上了校验反馈
        self.assertEqual(content["product_id"], 4643797)


class ContentToColumnsTests(unittest.TestCase):
    def test_maps_all_column_groups(self) -> None:
        content = _validate_material(_good_material(), 4643797,
                                     _assemble_fact_card(_full_row()))
        row = _content_to_columns(content)
        self.assertEqual(row["title_v1"], content["titles"][0]["title"])
        self.assertEqual(row["point_5"], content["short_points"][4])
        self.assertEqual(row["detail_pain"], "长时间行走脚部闷热易疲劳")
        self.assertIn("批发价：48 元", row["detail_specs"])
        self.assertEqual(row["brief_point"], "橡胶防滑大底纹路特写")


class ExportExcelTests(unittest.TestCase):
    def test_header_and_empty_audit_columns(self) -> None:
        content = _validate_material(_good_material(), 4643797,
                                     _assemble_fact_card(_full_row()))
        material_row = {
            "id": 7,
            "product_id": 4643797,
            "content_json": json.dumps(content, ensure_ascii=False),
            "machine_check": "pass",
            "risk_flags_json": json.dumps(["价格带偏低"], ensure_ascii=False),
        }
        with TemporaryDirectory() as tmp:
            with mock.patch.object(generator, "OUTPUT_DIR", Path(tmp)):
                path = _export_audit_excel(
                    [material_row], {4643797: _full_row()}, "20260101_000000_000"
                )
            self.assertTrue(path.exists())
            ws = load_workbook(path).active
            header = [ws.cell(row=1, column=c).value
                      for c in range(1, len(EXCEL_COLUMNS) + 1)]
            self.assertEqual(header, list(EXCEL_COLUMNS))
            values = {header[c - 1]: ws.cell(row=2, column=c).value
                      for c in range(1, len(header) + 1)}
            self.assertEqual(values["id"], 7)
            self.assertEqual(values["title_v1"], content["titles"][0]["title"])
            self.assertEqual(values["point_1"], content["short_points"][0])
            self.assertIsNone(values["audit_status"])
            self.assertIsNone(values["audit_operator"])
            self.assertIsNone(values["audit_comment"])


class _FakeCursor:
    def __init__(self, affected: int = 1) -> None:
        self.affected = affected
        self.sql = ""
        self.params = None

    def execute(self, sql, params=None):
        self.sql = sql
        self.params = params
        return self.affected


class _Ctx:
    def __init__(self, obj) -> None:
        self.obj = obj

    def __enter__(self):
        return self.obj

    def __exit__(self, *exc) -> bool:
        return False


class _FakeConn:
    def __init__(self, cur: _FakeCursor) -> None:
        self._cur = cur

    def cursor(self, *a, **kw):
        return _Ctx(self._cur)

    def __enter__(self):
        return self

    def __exit__(self, *exc) -> bool:
        return False


class MaterialAuditTests(unittest.TestCase):
    def _run_audit(self, current: dict | None, new_status: str,
                   comment: str = "", affected: int = 1):
        cur = _FakeCursor(affected=affected)
        with mock.patch.object(
            database, "fetch_one", return_value=current
        ), mock.patch.object(
            database, "connection", return_value=_FakeConn(cur)
        ):
            result = material_audit_product(
                current["id"] if current else 1, new_status, "tester", comment
            )
        return result, cur

    def test_pending_to_pass(self) -> None:
        result, cur = self._run_audit(
            {"id": 5, "audit_status": None}, "pass"
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["old_status"], "pending")
        self.assertIn("audit_status", cur.sql)
        self.assertEqual(cur.params[0], "pass")

    def test_reject_requires_comment(self) -> None:
        with mock.patch.object(
            database, "fetch_one", return_value={"id": 5, "audit_status": "pending"}
        ):
            with self.assertRaisesRegex(ValueError, "reject"):
                material_audit_product(5, "reject", "tester", "")

    def test_terminal_reject_blocks_reaudit(self) -> None:
        with mock.patch.object(
            database, "fetch_one", return_value={"id": 5, "audit_status": "reject"}
        ):
            with self.assertRaisesRegex(ValueError, "终态"):
                material_audit_product(5, "pass", "tester", "")

    def test_hold_to_pass_allowed(self) -> None:
        result, _ = self._run_audit({"id": 5, "audit_status": "hold"}, "pass")
        self.assertTrue(result["success"])

    def test_invalid_status_rejected(self) -> None:
        with mock.patch.object(
            database, "fetch_one", return_value={"id": 5, "audit_status": "pending"}
        ):
            with self.assertRaisesRegex(ValueError, "非法"):
                material_audit_product(5, "approved", "tester", "")


class ImportExcelTests(unittest.TestCase):
    @staticmethod
    def _write_excel(tmp: str, rows: list[list]) -> Path:
        path = Path(tmp) / "filled_material_audit.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "audit_status", "audit_operator", "audit_comment"])
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    def test_import_skips_blank_unfilled_and_collects_failures(self) -> None:
        # 行2: pass 正常；行3: 未填状态跳过；行4: reject 缺备注失败；行5: 全空行忽略
        with TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            path = self._write_excel(tmp, [
                [1, "pass", "alice", "文案合格"],
                [2, "", "", ""],
                [3, "reject", "bob", ""],
                [None, None, None, None],
            ])
            calls: list[tuple] = []

            def _fake_audit(material_id, status, operator, comment):
                calls.append((material_id, status, operator, comment))
                if status == "reject" and not comment:
                    raise ValueError("reject 状态必须填写审核备注原因")
                return {"success": True}

            with mock.patch.object(emi, "material_audit_product", side_effect=_fake_audit), \
                    mock.patch.object(emi, "_archive") as fake_archive, \
                    mock.patch("sys.stdout", new_callable=lambda: __import__("io").StringIO()):
                code = import_material_audit_excel(str(path))
            self.assertEqual(code, 0)
            self.assertEqual(calls, [
                (1, "pass", "alice", "文案合格"),
                (3, "reject", "bob", ""),
            ])
            fake_archive.assert_called_once()
            self.assertEqual(fake_archive.call_args[0][0], [1])  # 只归档成功行

    def test_missing_file_returns_2(self) -> None:
        self.assertEqual(import_material_audit_excel("no_such_file.xlsx"), 2)


class GenerateFlowTests(unittest.TestCase):
    def test_no_pass_products(self) -> None:
        with mock.patch.object(generator, "_ensure_table"), \
                mock.patch.object(generator, "_query_pass_product_ids", return_value=[]):
            result = generate_goods_material(export_excel=False)
        self.assertFalse(result["success"])
        self.assertEqual(result["count"], 0)
        self.assertIn("无审核通过", result["message"])

    def test_resume_skips_success_without_model_call(self) -> None:
        with mock.patch.object(generator, "_ensure_table"), \
                mock.patch.object(generator, "_query_pass_product_ids", return_value=[1]), \
                mock.patch.object(generator, "_query_info_rows", return_value={1: _full_row()}), \
                mock.patch.object(generator, "_query_material_rows", return_value={
                    1: {"product_id": 1, "gen_status": "success", "id": 9},
                }), \
                mock.patch.object(generator, "_call_qwen") as fake_call:
            result = generate_goods_material(export_excel=False)
        fake_call.assert_not_called()
        self.assertTrue(result["success"])
        self.assertEqual(result["generated"], 0)
        self.assertEqual(result["skipped_success"], 1)
        self.assertEqual(result["failed"], 0)

    def test_missing_info_marks_failed(self) -> None:
        with mock.patch.object(generator, "_ensure_table"), \
                mock.patch.object(generator, "_query_pass_product_ids", return_value=[2]), \
                mock.patch.object(generator, "_query_info_rows", return_value={}), \
                mock.patch.object(generator, "_query_material_rows", return_value={}), \
                mock.patch.object(generator, "_insert_pending_rows", return_value=1), \
                mock.patch.object(generator, "_mark_generating"), \
                mock.patch.object(generator, "_mark_failed") as fake_failed, \
                mock.patch.object(generator, "_mark_success") as fake_success:
            result = generate_goods_material(export_excel=False)
        fake_failed.assert_called_once_with(2, "product_info 缺失或标题为空")
        fake_success.assert_not_called()
        self.assertEqual(result["failed"], 1)
        self.assertFalse(result["success"])


class AgentToolMaterialTest(unittest.TestCase):
    """function_call_agent：generate_goods_mateial 工具注册 + 执行器注入 + wrapper 解析。"""

    def test_schema_registered_with_product_ids(self) -> None:
        from src.agent.function_call_agent import TOOLS

        tool = next(
            t["function"] for t in TOOLS
            if t["function"]["name"] == "generate_goods_mateial"
        )
        self.assertIn("product_ids", tool["parameters"]["properties"])
        self.assertEqual(tool["parameters"]["required"], ["product_ids"])
        self.assertIn("断点幂等", tool["description"])

    def test_system_prompt_documents_tool(self) -> None:
        from src.agent.function_call_agent import SYSTEM_PROMPT

        self.assertIn("generate_goods_mateial", SYSTEM_PROMPT)
        self.assertIn("audit_status='pass'", SYSTEM_PROMPT)

    def test_wrapper_parses_ids_and_calls_generator(self) -> None:
        from src.agent import function_call_agent as fca

        with mock.patch.object(fca, "generate_goods_material") as fake_gen:
            fake_gen.return_value = {"success": True, "generated": 1}
            result = fca.generate_goods_mateial(product_ids=[4643797, 123])
        fake_gen.assert_called_once_with(product_ids=[4643797, 123])
        self.assertTrue(result["success"])

    def test_wrapper_empty_list_falls_back_to_auto_query(self) -> None:
        from src.agent import function_call_agent as fca

        with mock.patch.object(fca, "generate_goods_material") as fake_gen:
            fake_gen.return_value = {"success": True, "generated": 0}
            fca.generate_goods_mateial(product_ids=[])
        fake_gen.assert_called_once_with(product_ids=None)

    def test_wrapper_rejects_bad_ids(self) -> None:
        from src.agent import function_call_agent as fca

        result = fca.generate_goods_mateial(product_ids=["abc"])
        self.assertFalse(result["success"])
        self.assertIn("非法", result["error"])

    def test_executor_injects_passed_product_ids(self) -> None:
        from src.agent import function_call_agent as fca

        state = {"passed_product_ids": [4643797]}

        def _fake_gen(product_ids=None, **kw):
            return {"success": True, "product_ids": product_ids}

        with mock.patch.object(fca, "generate_goods_material", _fake_gen):
            result = fca._execute_tool(
                "generate_goods_mateial", {"product_ids": []}, state
            )
        # 模型固定传 []，执行器注入 state 里的 pass 商品 ID
        self.assertEqual(result["product_ids"], [4643797])

    def test_executor_prefers_explicit_ids_over_state(self) -> None:
        from src.agent import function_call_agent as fca

        state = {"passed_product_ids": [1, 2, 3]}

        def _fake_gen(product_ids=None, **kw):
            return {"success": True, "product_ids": product_ids}

        with mock.patch.object(fca, "generate_goods_material", _fake_gen):
            result = fca._execute_tool(
                "generate_goods_mateial", {"product_ids": [99]}, state
            )
        self.assertEqual(result["product_ids"], [99])

    def test_executor_empty_state_passes_none_for_auto_query(self) -> None:
        from src.agent import function_call_agent as fca

        def _fake_gen(product_ids=None, **kw):
            return {"success": True, "product_ids": product_ids}

        with mock.patch.object(fca, "generate_goods_material", _fake_gen):
            result = fca._execute_tool(
                "generate_goods_mateial", {"product_ids": []}, {"passed_product_ids": []}
            )
        self.assertIsNone(result["product_ids"])

    def test_post_audit_chain_import_then_generate(self) -> None:
        """SYSTEM_PROMPT 审核后置链路：import 注入 pass ID -> generate 消费。"""
        from src.agent import function_call_agent as fca

        state = {"excel_path": None, "audit_pending": True, "passed_product_ids": []}

        def _fake_import(excel_file_path="", state=None):
            # 模拟 import_audit_excel 成功后回查 DB 注入 pass 商品 ID
            state["passed_product_ids"] = [4643797, 888]
            state["audit_pending"] = False
            return {"success": True, "passed_product_ids": [4643797, 888]}

        captured: list = []

        def _fake_gen(product_ids=None, **kw):
            captured.append(product_ids)
            return {"success": True, "generated": len(product_ids or [])}

        with mock.patch.object(fca, "import_audit_excel", _fake_import), \
                mock.patch.object(fca, "generate_goods_material", _fake_gen), \
                mock.patch.object(fca, "Path") as fake_path:
            fake_path.return_value.exists.return_value = True
            r1 = fca._execute_tool(
                "import_audit_excel", {"excel_file_path": "x.xlsx"}, state
            )
            r2 = fca._execute_tool(
                "generate_goods_mateial", {"product_ids": []}, state
            )
        self.assertTrue(r1["success"])
        self.assertTrue(r2["success"])
        # generate 拿到的是 import 刚注入的 pass 商品 ID
        self.assertEqual(captured, [[4643797, 888]])
        self.assertFalse(state["audit_pending"])

    def test_manual_audit_tool_still_refuses_model(self) -> None:
        """红线：manual_audit_product 仍拒绝模型自主审核。"""
        from src.agent import function_call_agent as fca

        result = fca._execute_tool(
            "manual_audit_product",
            {"product_ids": [1], "audit_status": "pass", "auditor": "gpt"},
            {},
        )
        self.assertFalse(result["success"])
        self.assertIn("仅人类调用", result["message"])


class AgentMaterialImportToolTest(unittest.TestCase):
    """素材审核导入工具：import_material_audit_excel 注册 + 分发路由 + state 注入。

    回归锁：修复前 Agent 只有选品导入工具，用户说“已审核素材表”会被误路由到
    import_audit_excel（选品表）；修复后必须走 import_material_audit_excel。
    """

    def test_schema_registered(self) -> None:
        from src.agent.function_call_agent import TOOLS

        tool = next(
            t["function"] for t in TOOLS
            if t["function"]["name"] == "import_material_audit_excel"
        )
        self.assertIn("excel_file_path", tool["parameters"]["properties"])
        self.assertEqual(tool["parameters"]["required"], ["excel_file_path"])
        self.assertIn("素材审核Excel", tool["description"])

    def test_system_prompt_documents_tool_and_disambiguates(self) -> None:
        from src.agent.function_call_agent import SYSTEM_PROMPT

        self.assertIn("import_material_audit_excel", SYSTEM_PROMPT)
        # 两条审核链路按触发词消歧：素材表 -> import_material_audit_excel
        self.assertIn("已审核素材表", SYSTEM_PROMPT)
        # 素材链路禁止误用选品导入工具
        self.assertIn("禁止**调用 import_audit_excel", SYSTEM_PROMPT)

    def test_executor_routes_material_import_not_selection(self) -> None:
        """分发：import_material_audit_excel 走素材导入 wrapper（含 state 注入），
        不得落到选品导入 import_audit_excel。"""
        from src.agent import function_call_agent as fca

        state = {"material_excel_path": "m.xlsx"}

        def _fake_run(path):
            print("[汇总] 成功更新 1 条")
            print("[汇总] 跳过 0 条未填写审核状态")
            print("[汇总] 失败 0 条")
            return 0

        selection_called: list = []

        def _fake_sel_import(excel_file_path="", state=None):
            selection_called.append(excel_file_path)
            return {"success": True, "via": "selection"}

        # 只 mock 底层导入函数 + Path，真实 wrapper 执行（验证其 state 注入）
        with mock.patch.object(fca, "_run_material_import", _fake_run), \
                mock.patch.object(fca, "Path") as fake_path, \
                mock.patch.object(fca, "import_audit_excel", _fake_sel_import):
            fake_path.return_value.exists.return_value = True
            result = fca._execute_tool(
                "import_material_audit_excel", {"excel_file_path": ""}, state
            )
        self.assertTrue(result["success"])
        # 空路径 -> 真实 wrapper 从 state["material_excel_path"] 注入
        self.assertEqual(result["excel_path"], "m.xlsx")
        # 选品导入工具绝不能被调用
        self.assertEqual(selection_called, [])

    def test_generate_stores_material_excel_path_in_state(self) -> None:
        """generate_goods_mateial 返回的 excel_path 应存入 state['material_excel_path']。"""
        from src.agent import function_call_agent as fca

        state = {"passed_product_ids": [1]}

        def _fake_gen(product_ids=None, **kw):
            return {"success": True, "excel_path": "out/pending_material_audit.xlsx"}

        with mock.patch.object(fca, "generate_goods_material", _fake_gen):
            fca._execute_tool("generate_goods_mateial", {"product_ids": []}, state)
        self.assertEqual(state["material_excel_path"], "out/pending_material_audit.xlsx")
        self.assertTrue(state["material_audit_pending"])

    def test_material_import_missing_state_path_errors(self) -> None:
        """state 无素材 Excel 路径且未传路径 -> 返回 error，不误用选品路径。"""
        from src.agent import function_call_agent as fca

        result = fca._execute_tool(
            "import_material_audit_excel",
            {"excel_file_path": ""},
            {"excel_path": "selection.xlsx"},  # 只有选品路径，不应被采用
        )
        self.assertFalse(result["success"])
        self.assertIn("material_excel_path", result["error"])

    def test_material_import_wrapper_parses_stats(self) -> None:
        """wrapper 调真实导入函数并解析统计（stdout 汇总行）。"""
        from src.agent import function_call_agent as fca

        def _fake_run(path):
            print("[汇总] 成功更新 2 条")
            print("[汇总] 跳过 1 条未填写审核状态")
            print("[汇总] 失败 0 条")
            return 0

        state = {"material_excel_path": "m.xlsx"}
        with mock.patch.object(fca, "Path") as fake_path, \
                mock.patch.object(fca, "_run_material_import", _fake_run):
            fake_path.return_value.exists.return_value = True
            result = fca.import_material_audit_excel("", state=state)
        self.assertTrue(result["success"])
        self.assertEqual(result["success_updated"], 2)
        self.assertEqual(result["skipped_not_filled"], 1)
        self.assertEqual(result["failed"], 0)
        self.assertFalse(state["material_audit_pending"])


if __name__ == "__main__":
    unittest.main()
